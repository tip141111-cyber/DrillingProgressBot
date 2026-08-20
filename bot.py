from __future__ import annotations
import json
import logging
import os
from datetime import datetime, time
from pathlib import Path
try:
    from zoneinfo import ZoneInfo
except ModuleNotFoundError:
    from backports.zoneinfo import ZoneInfo

from dotenv import load_dotenv
from openpyxl import load_workbook
from telegram import ReplyKeyboardMarkup, Update
from telegram.constants import ChatAction
from telegram.ext import (
    Application,
    ApplicationBuilder,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_EXCEL_PATH = BASE_DIR / "data" / "uchet_bureniya_otverstiy.xlsx"
CHAT_STORE_PATH = BASE_DIR / "chats.json"
DEFAULT_TIMEZONE = "Asia/Yakutsk"
BUTTON_RECORD = "Записать"
BUTTON_CORRECT = "Откорректировать"
BUTTON_CANCEL = "Отмена"
BUTTON_DOWNLOAD = "Скачать актуальную"
BUTTON_AFU_14 = "АФУ 1/4"
BUTTON_AFU_3 = "АФУ 3"
BUTTON_BOTH_LINES = "Обе линейки"
MAIN_KEYBOARD = ReplyKeyboardMarkup(
    [[BUTTON_RECORD, BUTTON_CORRECT], [BUTTON_DOWNLOAD, BUTTON_CANCEL]],
    resize_keyboard=True,
)
LINE_KEYBOARD = ReplyKeyboardMarkup(
    [[BUTTON_AFU_14, BUTTON_AFU_3], [BUTTON_BOTH_LINES], [BUTTON_CANCEL]],
    resize_keyboard=True,
)

QUESTIONS = [
    ("АФУ 1/4", "разметка", "C8"),
    ("АФУ 1/4", "бурение", "D8"),
    ("АФУ 1/4", "гофра", "E8"),
    ("АФУ 3", "разметка", "C9"),
    ("АФУ 3", "бурение", "D9"),
    ("АФУ 3", "гофра", "E9"),
]
LINE_QUESTIONS = {
    BUTTON_AFU_14: QUESTIONS[:3],
    BUTTON_AFU_3: QUESTIONS[3:],
    BUTTON_BOTH_LINES: QUESTIONS,
}

logging.basicConfig(
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)


def get_excel_path() -> Path:
    return Path(os.getenv("EXCEL_PATH", str(DEFAULT_EXCEL_PATH))).expanduser().resolve()


def get_timezone() -> str:
    return os.getenv("BOT_TIMEZONE", DEFAULT_TIMEZONE)


def get_prompt_time() -> time:
    hour = int(os.getenv("BOT_PROMPT_HOUR", "17"))
    minute = int(os.getenv("BOT_PROMPT_MINUTE", "10"))
    return time(hour=hour, minute=minute, tzinfo=ZoneInfo(get_timezone()))


def get_admin_chat_ids() -> set[int]:
    raw_value = os.getenv("ADMIN_CHAT_IDS", "")
    admin_ids = set()
    for item in raw_value.split(","):
        item = item.strip()
        if not item:
            continue
        try:
            admin_ids.add(int(item))
        except ValueError:
            logger.warning("Некорректный ADMIN_CHAT_IDS: %s", item)
    return admin_ids


def is_admin_chat(chat_id: int) -> bool:
    return chat_id in get_admin_chat_ids()


async def deny_access(update: Update) -> None:
    chat_id = update.effective_chat.id
    await update.message.reply_text(
        "Нет доступа к этому боту.\n"
        f"Ваш chat_id: {chat_id}\n\n"
        "Если это твой второй аккаунт или другой чат, добавь этот ID в ADMIN_CHAT_IDS."
    )


def load_chat_ids() -> set[int]:
    return {int(chat_id) for chat_id in load_store().get("chat_ids", [])}


def load_store() -> dict:
    if not CHAT_STORE_PATH.exists():
        return {"chat_ids": [], "submitted_dates": {}}
    try:
        store = json.loads(CHAT_STORE_PATH.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {"chat_ids": [], "submitted_dates": {}}
    store.setdefault("chat_ids", [])
    store.setdefault("submitted_dates", {})
    return store


def save_store(store: dict) -> None:
    CHAT_STORE_PATH.write_text(
        json.dumps(store, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def save_chat_id(chat_id: int) -> None:
    store = load_store()
    chat_ids = {int(saved_chat_id) for saved_chat_id in store.get("chat_ids", [])}
    chat_ids.add(chat_id)
    store["chat_ids"] = sorted(chat_ids)
    save_store(store)


def today_key() -> str:
    return datetime.now(ZoneInfo(get_timezone())).strftime("%Y-%m-%d")


def mark_submitted_today(chat_id: int) -> None:
    store = load_store()
    store.setdefault("submitted_dates", {})[str(chat_id)] = today_key()
    save_store(store)


def has_submitted_today(chat_id: int) -> bool:
    store = load_store()
    return store.get("submitted_dates", {}).get(str(chat_id)) == today_key()


def get_sessions(context: ContextTypes.DEFAULT_TYPE) -> dict:
    return context.application.bot_data.setdefault("sessions", {})


def get_session(chat_id: int, context: ContextTypes.DEFAULT_TYPE) -> dict | None:
    return get_sessions(context).get(chat_id)


def start_session(chat_id: int, context: ContextTypes.DEFAULT_TYPE, mode: str = "add") -> dict:
    session = {"step": 0, "answers": {}, "mode": mode, "questions": [], "awaiting_line": True}
    get_sessions(context)[chat_id] = session
    return session


def clear_session(chat_id: int, context: ContextTypes.DEFAULT_TYPE) -> None:
    get_sessions(context).pop(chat_id, None)


def build_line_choice(mode: str = "add") -> str:
    action = "записать" if mode == "add" else "откорректировать"
    return f"Выбери линейку, по которой нужно {action} значения."


def build_question(step: int, mode: str = "add", questions: list[tuple[str, str, str]] | None = None) -> str:
    questions = questions or QUESTIONS
    line_name, stage_name, _cell = questions[step]
    action = "сколько добавить за сегодня" if mode == "add" else "какое значение поставить вручную"
    return (
        f"{step + 1}/{len(questions)}. {line_name}: {action} "
        f"по этапу «{stage_name}»? Введите число от 0 до 144."
    )


def validate_count(text: str) -> int | None:
    normalized = text.strip().replace(",", ".")
    try:
        value = float(normalized)
    except ValueError:
        return None
    if not value.is_integer():
        return None
    value = int(value)
    if value < 0 or value > 144:
        return None
    return value


def update_workbook(answers: dict[str, int], mode: str = "add") -> tuple[Path, dict[str, float]]:
    excel_path = get_excel_path()
    workbook = load_workbook(excel_path)
    sheet = workbook["Учет работ"] if "Учет работ" in workbook.sheetnames else workbook.active

    for cell_address, value in answers.items():
        if mode == "set":
            sheet[cell_address] = value
        else:
            current_value = number(sheet[cell_address].value)
            sheet[cell_address] = min(144, current_value + value)

    total_plan = number(sheet["B8"].value) + number(sheet["B9"].value)
    total_marked = number(sheet["C8"].value) + number(sheet["C9"].value)
    total_drilled = number(sheet["D8"].value) + number(sheet["D9"].value)
    total_corrugated = number(sheet["E8"].value) + number(sheet["E9"].value)

    sheet["B10"] = total_plan
    sheet["C10"] = total_marked
    sheet["D10"] = total_drilled
    sheet["E10"] = total_corrugated

    afu_14_progress = calculate_progress(sheet["B8"].value, sheet["C8"].value, sheet["D8"].value, sheet["E8"].value)
    afu_3_progress = calculate_progress(sheet["B9"].value, sheet["C9"].value, sheet["D9"].value, sheet["E9"].value)
    total_progress = calculate_progress(total_plan, total_marked, total_drilled, total_corrugated)

    sheet["F8"] = afu_14_progress
    sheet["F9"] = afu_3_progress
    sheet["F10"] = total_progress
    sheet["I5"] = total_progress

    sheet["I8"] = afu_14_progress * 100
    sheet["I9"] = afu_3_progress * 100
    sheet["I10"] = total_progress * 100
    apply_number_formats(sheet)

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        pass

    try:
        workbook.save(excel_path)
        return excel_path, read_current_values(sheet)
    except PermissionError:
        updated_dir = BASE_DIR / "updated"
        updated_dir.mkdir(exist_ok=True)
        timestamp = datetime.now(ZoneInfo(get_timezone())).strftime("%Y%m%d_%H%M")
        fallback_path = updated_dir / f"uchet_bureniya_otverstiy_{timestamp}.xlsx"
        workbook.save(fallback_path)
        return fallback_path, read_current_values(sheet)


def apply_number_formats(sheet) -> None:
    for cell_address in ("F8", "F9", "F10", "I5"):
        sheet[cell_address].number_format = "0.0%"
    for cell_address in ("I8", "I9", "I10"):
        sheet[cell_address].number_format = '0.0'


def read_current_values(sheet) -> dict[str, float]:
    return {
        "C8": number(sheet["C8"].value),
        "D8": number(sheet["D8"].value),
        "E8": number(sheet["E8"].value),
        "C9": number(sheet["C9"].value),
        "D9": number(sheet["D9"].value),
        "E9": number(sheet["E9"].value),
    }


def calculate_progress(plan: int | float, marked: int | float, drilled: int | float, corrugated: int | float) -> float:
    plan = number(plan)
    if plan <= 0:
        return 0.0
    marked = number(marked)
    drilled = number(drilled)
    corrugated = number(corrugated)
    return min(1.0, max(0.0, (marked * 0.1 + drilled * 0.7 + corrugated * 0.2) / plan))


def number(value: object) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    return 0.0


def build_warning(answers: dict[str, int]) -> str:
    warnings = []
    for line_name, marked, drilled, corrugated in (
        ("АФУ 1/4", answers.get("C8", 0), answers.get("D8", 0), answers.get("E8", 0)),
        ("АФУ 3", answers.get("C9", 0), answers.get("D9", 0), answers.get("E9", 0)),
    ):
        if drilled > marked:
            warnings.append(f"{line_name}: бурения больше, чем разметки.")
        if corrugated > drilled:
            warnings.append(f"{line_name}: гофры больше, чем бурения.")

    if not warnings:
        return ""
    return "\n\nПроверь логику чисел:\n" + "\n".join(f"- {item}" for item in warnings)


async def send_workbook(chat_id: int, context: ContextTypes.DEFAULT_TYPE, path: Path) -> None:
    await context.bot.send_chat_action(chat_id=chat_id, action=ChatAction.UPLOAD_DOCUMENT)
    with path.open("rb") as file:
        await context.bot.send_document(
            chat_id=chat_id,
            document=file,
            filename=path.name,
            caption="Актуальная таблица учета работ.",
        )


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    chat_id = update.effective_chat.id
    if not is_admin_chat(chat_id):
        await deny_access(update)
        return
    save_chat_id(chat_id)
    await update.message.reply_text(
        "Готов. Я буду спрашивать данные по будням в 17:10, если ты сам не внес отчет за день.\n\n"
        "Кнопки:\n"
        "Записать - добавить данные за день\n"
        "Откорректировать - вручную поставить точные значения\n"
        "Скачать актуальную - получить таблицу\n"
        "Отмена - отменить текущий ввод",
        reply_markup=MAIN_KEYBOARD,
    )


async def begin_report(chat_id: int, context: ContextTypes.DEFAULT_TYPE, mode: str = "add") -> str:
    start_session(chat_id, context, mode)
    if mode == "set":
        return "Режим корректировки: введенные числа заменят текущие значения.\n" + build_line_choice(mode)
    return "Режим записи: введенные числа прибавятся к текущим значениям.\n" + build_line_choice(mode)


async def report(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    chat_id = update.effective_chat.id
    if not is_admin_chat(chat_id):
        await deny_access(update)
        return
    save_chat_id(chat_id)
    await update.message.reply_text(await begin_report(chat_id, context, "add"), reply_markup=LINE_KEYBOARD)


async def download(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    chat_id = update.effective_chat.id
    if not is_admin_chat(chat_id):
        await deny_access(update)
        return
    path = get_excel_path()
    if not path.exists():
        await update.message.reply_text(f"Не нашел файл таблицы: {path}")
        return
    await send_workbook(chat_id, context, path)


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not is_admin_chat(update.effective_chat.id):
        await deny_access(update)
        return
    clear_session(update.effective_chat.id, context)
    await update.message.reply_text("Ок, текущий ввод отменил.", reply_markup=MAIN_KEYBOARD)


async def handle_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    chat_id = update.effective_chat.id
    if not is_admin_chat(chat_id):
        await deny_access(update)
        return
    text = update.message.text.strip()

    if text == BUTTON_RECORD:
        save_chat_id(chat_id)
        await update.message.reply_text(await begin_report(chat_id, context, "add"), reply_markup=LINE_KEYBOARD)
        return
    if text == BUTTON_CORRECT:
        save_chat_id(chat_id)
        await update.message.reply_text(await begin_report(chat_id, context, "set"), reply_markup=LINE_KEYBOARD)
        return
    if text == BUTTON_DOWNLOAD:
        await download(update, context)
        return
    if text == BUTTON_CANCEL:
        await cancel(update, context)
        return

    session = get_session(chat_id, context)
    if not session:
        await update.message.reply_text("Чтобы внести данные, нажми кнопку «Записать».", reply_markup=MAIN_KEYBOARD)
        return

    if session.get("awaiting_line"):
        if text not in LINE_QUESTIONS:
            await update.message.reply_text("Выбери линейку кнопкой ниже.", reply_markup=LINE_KEYBOARD)
            return
        session["questions"] = LINE_QUESTIONS[text]
        session["awaiting_line"] = False
        session["step"] = 0
        await update.message.reply_text(
            build_question(0, session.get("mode", "add"), session["questions"]),
            reply_markup=MAIN_KEYBOARD,
        )
        return

    value = validate_count(text)
    if value is None:
        await update.message.reply_text("Введите целое число от 0 до 144.", reply_markup=MAIN_KEYBOARD)
        return

    step = session["step"]
    mode = session.get("mode", "add")
    questions = session.get("questions") or QUESTIONS
    _line_name, _stage_name, cell_address = questions[step]
    session["answers"][cell_address] = value
    session["step"] += 1

    if session["step"] < len(questions):
        await update.message.reply_text(build_question(session["step"], mode, questions), reply_markup=MAIN_KEYBOARD)
        return

    answers = session["answers"]
    clear_session(chat_id, context)
    path, current_values = update_workbook(answers, mode)
    mark_submitted_today(chat_id)
    warning = build_warning(current_values)
    result_text = "Данные добавлены в таблицу." if mode == "add" else "Значения откорректированы вручную."
    await update.message.reply_text(result_text + warning, reply_markup=MAIN_KEYBOARD)
    await send_workbook(chat_id, context, path)


async def daily_prompt(context: ContextTypes.DEFAULT_TYPE) -> None:
    now = datetime.now(ZoneInfo(get_timezone()))
    if now.weekday() >= 5:
        return
    for chat_id in load_chat_ids():
        if not is_admin_chat(chat_id):
            continue
        if has_submitted_today(chat_id):
            continue
        start_session(chat_id, context, "add")
        await context.bot.send_message(
            chat_id=chat_id,
            text="Время ежедневного отчета по отверстиям.\n" + build_line_choice("add"),
            reply_markup=LINE_KEYBOARD,
        )


def main() -> None:
    load_dotenv(BASE_DIR / ".env")
    token = os.getenv("TELEGRAM_BOT_TOKEN")
    if not token:
        raise RuntimeError("Укажите TELEGRAM_BOT_TOKEN в .env или переменных окружения.")
    if not get_admin_chat_ids():
        raise RuntimeError("Укажите ADMIN_CHAT_IDS в .env. Узнать ID можно из chats.json или сообщения /start.")

    application: Application = ApplicationBuilder().token(token).build()
    application.add_handler(CommandHandler("start", start))
    application.add_handler(CommandHandler("report", report))
    application.add_handler(CommandHandler("download", download))
    application.add_handler(CommandHandler("cancel", cancel))
    application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_answer))

    application.job_queue.run_daily(
        daily_prompt,
        time=get_prompt_time(),
    )
    application.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()

